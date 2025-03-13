import streamlit as st
from fpdf import FPDF  # For PDF generation
from docx import Document  # For Word document generation
import re  # For sanitizing file names
from cover_letter import generate_cover_letter  # Import your cover letter generator function

def generate_pdf(text):
    """
    Converts the given text to a PDF using the FPDF library and returns the binary data.
    """
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    # Add a centered, bold title
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Cover Letter", ln=True, align="C")
    pdf.ln(10)  # Add some space after the title

    # Add the cover letter text
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, text.encode('latin1', 'replace').decode('latin1'))

    return pdf.output(dest='S').encode('latin1')

def generate_docx(text):
    """
    Converts the given text to a Word document using the python-docx library and returns the binary data.
    """
    doc = Document()
    doc.add_paragraph(text)
    temp_file_path = "/tmp/cover_letter.docx"
    doc.save(temp_file_path)
    with open(temp_file_path, 'rb') as file:
        docx_data = file.read()
    return docx_data

def cover_letter_section():
    st.markdown("### ✉️ Cover Letter Details")
    st.markdown("---")
    
    # Initialize session state for cover letter data
    if 'cover_letter_info' not in st.session_state:
        st.session_state.cover_letter_info = {}

    # Cover Letter Form
    applicant_name = st.text_input(
        "Your Name (Required):",
        placeholder="e.g., John Doe",
        value=st.session_state.cover_letter_info.get('name', '')
    )
    hiring_manager = st.text_input(
        "Hiring Manager Name (Optional):",
        placeholder="e.g., Jane Smith",
        value=st.session_state.cover_letter_info.get('hiring_manager', '')
    )
    job_title = st.text_input(
        "Job Title You Are Applying For (Required):",
        placeholder="e.g., Data Analyst",
        value=st.session_state.cover_letter_info.get('job_title', '')
    )
    company_name = st.text_input(
        "Company Name (Required):",
        placeholder="e.g., TechCorp Inc.",
        value=st.session_state.cover_letter_info.get('company_name', '')
    )
    job_description = st.text_area(
        "Job Description (Paste the job description here) (Required):",
        placeholder="E.g., We're looking for a highly skilled Data Analyst with expertise in Python and SQL...",
        value=st.session_state.cover_letter_info.get('job_description', '')
    )

    # Auto-save inputs to session state as the user types
    st.session_state.cover_letter_info.update({
        'name': applicant_name,
        'hiring_manager': hiring_manager,
        'job_title': job_title,
        'company_name': company_name,
        'job_description': job_description
    })

    # Check if all required fields are filled before generating cover letter
    required_fields = [applicant_name, job_title, company_name, job_description]
    if not all(required_fields):
        st.error("Please fill in all required fields (Name, Job Title, Company Name, and Job Description).")

    # Generate Cover Letter Button
    generate_button = st.button("Generate Cover Letter", disabled=not all(required_fields))

    if generate_button:
        try:
            # Debug: Print the data being passed to the function
            st.write("Data being passed to generate_cover_letter:", st.session_state.cover_letter_info)

            # Generate the cover letter text using the backend function
            cover_letter = generate_cover_letter(st.session_state.cover_letter_info)

            if not cover_letter:
                st.error("Failed to generate cover letter. Please try again.")
                return

            # Personalize greeting if the hiring manager's name is provided
            if hiring_manager:
                cover_letter = cover_letter.replace("Dear Hiring Manager", f"Dear {hiring_manager}")

            # Display Generated Cover Letter
            st.text_area("Generated Cover Letter:", value=cover_letter, height=300)
            st.success("Cover Letter Generated!")

            # Download Options for PDF and DOCX
            download_format = st.radio("Download Format", options=["PDF (.pdf)", "Word (.docx)"], index=0)

            # Sanitize job title for file names (avoid invalid characters)
            safe_job_title = re.sub(r'[\\/*?:"<>|]', "_", job_title)

            if download_format == "PDF (.pdf)":
                pdf_data = generate_pdf(cover_letter)
                st.download_button(
                    label="Download Cover Letter as PDF",
                    data=pdf_data,
                    file_name=f"{safe_job_title}_Cover_Letter.pdf",
                    mime="application/pdf"
                )
            elif download_format == "Word (.docx)":
                docx_data = generate_docx(cover_letter)
                st.download_button(
                    label="Download Cover Letter as Word Document",
                    data=docx_data,
                    file_name=f"{safe_job_title}_Cover_Letter.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

        except Exception as e:
            st.error(f"Error generating cover letter: {e}")

# Run the Streamlit app
if __name__ == "__main__":
    cover_letter_section()
